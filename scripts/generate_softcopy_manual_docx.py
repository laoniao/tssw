# -*- coding: utf-8 -*-
"""
生成软著登记用《软件说明书》Word 文档（文档鉴别材料）

排版要点（符合常见软著交存要求且兼顾美观）：
- A4，页眉标注软件名称与版本，页脚页码
- 正文宋体小四，固定行距约 18–20 磅，保证打印每页不少于 30 行
- 标题黑体分级，正文首行缩进 2 字符，由 Word 自然换行（禁止生硬截断）
- 全文不足 60 页提交全部；达到或超过 60 页则前、后各连续 30 页
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from softcopy_manual_content import SOFTWARE_NAME, SOFTWARE_VERSION, build_manual_blocks

# 软著：文档每页不少于 30 行
LINES_PER_PAGE = 30
FRONT_PAGES = 30
BACK_PAGES = 30
# 估算：正文区约可容纳字数（用于是否截取前后 30 页）
CHARS_PER_LINE = 38
EST_LINES_PER_PAGE = 32

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / f"{SOFTWARE_NAME}_{SOFTWARE_VERSION}_软件说明书.docx"


def set_run_font(run, *, cn="宋体", en="Times New Roman", size_pt=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def set_line_spacing(paragraph, pt=18):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(pt)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, cn="宋体", en="宋体", size_pt=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def set_section_page_start(section, start: int):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:start"), str(start))


def setup_section(section, page_start: int = 1):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # 常规公文页边距；配合 18 磅行距，正文区约 32–34 行（≥30）
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    set_section_page_start(section, page_start)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.clear()
    pf = hp.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    set_line_spacing(hp, 14)
    run = hp.add_run(f"{SOFTWARE_NAME} {SOFTWARE_VERSION}")
    set_run_font(run, cn="宋体", en="宋体", size_pt=10.5)

    # 页眉下划线
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "666666")
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.clear()
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)
    set_line_spacing(fp, 14)
    left = fp.add_run("第 ")
    set_run_font(left, cn="宋体", en="宋体", size_pt=9)
    add_page_field(fp)
    right = fp.add_run(" 页")
    set_run_font(right, cn="宋体", en="宋体", size_pt=9)


def add_blank(doc: Document, count: int = 1):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        set_line_spacing(p, 18)
        run = p.add_run(" ")
        set_run_font(run, size_pt=12)


def add_cover(doc: Document, blocks: list[tuple[str, str]]):
    """封面：居中标题 + 元信息。"""
    cover_items = [b for b in blocks if b[0].startswith("cover_")]
    cover_note = next(
        (t for k, t in blocks if k == "note" and t.startswith("本说明书用于")),
        "",
    )
    # 上方留白
    add_blank(doc, 6)

    for kind, text in cover_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        set_line_spacing(p, 22)

        if kind == "cover_title":
            pf.space_before = Pt(8)
            pf.space_after = Pt(8)
            set_line_spacing(p, 28)
            run = p.add_run(text)
            set_run_font(run, cn="黑体", en="Arial", size_pt=22, bold=True)
        elif kind == "cover_sub":
            pf.space_after = Pt(6)
            run = p.add_run(text)
            set_run_font(run, cn="黑体", en="Arial", size_pt=16, bold=True)
        elif kind == "cover_meta":
            set_line_spacing(p, 20)
            run = p.add_run(text)
            set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)

    if cover_note:
        add_blank(doc, 2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(0)
        set_line_spacing(p, 18)
        run = p.add_run(cover_note)
        set_run_font(run, cn="宋体", en="Times New Roman", size_pt=10.5)

    # 封面后分页进入正文
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8 if level == 2 else 6)
    pf.space_after = Pt(6 if level == 1 else 4)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(22 if level == 1 else 20)
    pf.keep_with_next = True

    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, cn="黑体", en="Arial", size_pt=16, bold=True)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        set_run_font(run, cn="黑体", en="Arial", size_pt=14, bold=True)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        set_run_font(run, cn="黑体", en="Arial", size_pt=12, bold=True)


def add_body(doc: Document, text: str, *, first_indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_line_spacing(p, 18)
    if first_indent:
        pf.first_line_indent = Cm(0.74)  # 约 2 字符
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)


def add_bullet(doc: Document, text: str, marker="·"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_line_spacing(p, 18)
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.37)
    run = p.add_run(f"{marker} {text}")
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)


def add_step(doc: Document, text: str, index: int):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_line_spacing(p, 18)
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.37)
    run = p.add_run(f"{index}. {text}")
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)


def add_toc_item(doc: Document, text: str):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_line_spacing(p, 18)
    # 保留前导空格缩进
    indent_level = (len(text) - len(text.lstrip(" "))) // 4
    pf.left_indent = Cm(0.5 * indent_level)
    run = p.add_run(text.strip())
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)


def add_note(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    set_line_spacing(p, 18)
    run = p.add_run(text)
    set_run_font(run, cn="楷体", en="Times New Roman", size_pt=12)


def estimate_lines(blocks: list[tuple[str, str]]) -> int:
    total = 0
    for kind, text in blocks:
        if kind.startswith("cover_"):
            total += 2
        elif kind == "h1":
            total += 3
        elif kind == "h2":
            total += 2
        elif kind == "h3":
            total += 2
        elif kind in ("p", "note"):
            total += max(1, math.ceil(len(text) / CHARS_PER_LINE))
        elif kind in ("bullet", "step", "toc"):
            total += max(1, math.ceil((len(text) + 3) / CHARS_PER_LINE))
        else:
            total += 1
    return total


def write_body(doc: Document, blocks: list[tuple[str, str]]):
    step_idx = 0
    for kind, text in blocks:
        if kind.startswith("cover_"):
            continue
        if kind == "note" and text.startswith("本说明书用于"):
            # 封面 note 已写入封面
            continue

        if kind == "h1":
            step_idx = 0
            add_heading(doc, text, 1)
        elif kind == "h2":
            step_idx = 0
            add_heading(doc, text, 2)
        elif kind == "h3":
            step_idx = 0
            add_heading(doc, text, 3)
        elif kind == "p":
            add_body(doc, text, first_indent=True)
        elif kind == "bullet":
            add_bullet(doc, text)
        elif kind == "step":
            step_idx += 1
            add_step(doc, text, step_idx)
        elif kind == "toc":
            add_toc_item(doc, text)
        elif kind == "note":
            add_note(doc, text)


def split_blocks_for_deposit(blocks: list[tuple[str, str]]):
    """不足 60 页返回全部；否则按估算行数切前 30 页 + 后 30 页对应内容块。"""
    total_lines = estimate_lines(blocks)
    total_pages = max(1, math.ceil(total_lines / EST_LINES_PER_PAGE))
    if total_pages < FRONT_PAGES + BACK_PAGES:
        return blocks, [], total_pages

    front_budget = FRONT_PAGES * EST_LINES_PER_PAGE
    back_budget = BACK_PAGES * EST_LINES_PER_PAGE

    front: list[tuple[str, str]] = []
    used = 0
    for b in blocks:
        need = estimate_lines([b])
        if used + need > front_budget and front:
            break
        front.append(b)
        used += need

    back: list[tuple[str, str]] = []
    used = 0
    for b in reversed(blocks):
        need = estimate_lines([b])
        if used + need > back_budget and back:
            break
        back.append(b)
        used += need
    back.reverse()
    return front, back, total_pages


def clear_default_empty(doc: Document):
    if doc.paragraphs and not doc.paragraphs[0].text.strip():
        el = doc.paragraphs[0]._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def main():
    blocks = build_manual_blocks()
    front, back, total_pages = split_blocks_for_deposit(blocks)

    doc = Document()
    # 默认 Normal 样式
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    setup_section(doc.sections[0], 1)

    # 封面 + 正文（不足 60 页交全文；超过则前后各 30 页）
    add_cover(doc, blocks if back else front)
    body_source = front
    body_blocks = [
        b for b in body_source
        if not b[0].startswith("cover_")
        and not (b[0] == "note" and b[1].startswith("本说明书用于"))
    ]
    write_body(doc, body_blocks)
    clear_default_empty(doc)

    if back:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        # 后 30 页页码从估算总页数往前推
        start = max(1, total_pages - BACK_PAGES + 1)
        setup_section(section, start)
        write_body(doc, back)

    doc.save(OUT_PATH)

    exported = total_pages if not back else FRONT_PAGES + BACK_PAGES
    print(f"软件名称: {SOFTWARE_NAME} {SOFTWARE_VERSION}")
    print(f"文档类型: 软件说明书")
    print(f"估算总页数: {total_pages}（目标每页≥{LINES_PER_PAGE}行）")
    if back:
        print(f"交存: 前{FRONT_PAGES}页 + 后{BACK_PAGES}页")
    else:
        print(f"交存: 全部约 {exported} 页（不足60页，提交整个文档）")
    print(f"输出文件: {OUT_PATH}")


if __name__ == "__main__":
    main()
